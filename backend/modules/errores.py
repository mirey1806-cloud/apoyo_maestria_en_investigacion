"""
Módulo 3: Detector de errores metodológicos.

Sigue el mismo patrón que los módulos 1 y 2:
- Requiere sesión iniciada (login)
- Soporta varios idiomas
- Permite pegar texto y/o adjuntar un PDF
- Permite exportar el reporte a Word

Evalúa la COHERENCIA INTERNA entre objetivos, hipótesis, metodología y
conclusiones de un documento académico (no ortografía ni formato).
"""
import io
import re
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document as DocxDocument

from backend.core.llm_client import generar_respuesta
from backend.core.config import IDIOMAS_SOPORTADOS, TAMANO_MAXIMO_ARCHIVO_MB
from backend.modules.auth import usuario_autenticado

router = APIRouter(prefix="/api/errores", tags=["Detector de errores metodológicos"])

SYSTEM_PROMPT = """Eres un revisor metodológico experto en investigación científica.
Analiza el documento proporcionado y evalúa ÚNICAMENTE la coherencia metodológica
entre estos cuatro elementos: objetivos, hipótesis (si existen), metodología y
conclusiones/resultados.

Para cada inconsistencia encontrada, entrega:
- Elementos en conflicto (por ejemplo: Objetivo vs. Metodología)
- Cita breve (máximo 20 palabras) de cada parte en conflicto
- Explicación clara de por qué son inconsistentes
- Sugerencia concreta de corrección

Si algún elemento no está presente en el documento (por ejemplo, no hay
hipótesis), indícalo como "elemento ausente" en vez de inventar una inconsistencia.

Al final entrega un resumen con:
- Nivel de coherencia general (Alto / Medio / Bajo)
- Lista priorizada de las 3 correcciones más urgentes

Reglas:
- No corrijas redacción, ortografía ni formato APA: eso está fuera de tu alcance.
- No inventes contenido que no esté en el documento.
- Usa formato Markdown simple (títulos con ###, listas con guiones).
- Responde en el idioma que te indique el usuario.
"""


class SolicitudRevision(BaseModel):
    idioma: str = Field(default="es")
    nombre_documento: str | None = Field(default=None)
    texto_documento: str = Field(..., description="Texto del documento a revisar")


class ReporteErrores(BaseModel):
    reporte_markdown: str


def _validar_idioma(idioma: str) -> None:
    if idioma not in IDIOMAS_SOPORTADOS:
        raise HTTPException(
            status_code=422,
            detail=f"Idioma '{idioma}' no soportado. Usa uno de: {', '.join(IDIOMAS_SOPORTADOS)}",
        )


def _mock_reporte(nombre: str | None) -> str:
    return f"""### Documento analizado: {nombre or 'sin nombre'} (resultado de ejemplo)

### Inconsistencia 1: Objetivo vs. Metodología
- **Cita del objetivo:** "[fragmento de ejemplo del objetivo]"
- **Cita de la metodología:** "[fragmento de ejemplo de la metodología]"
- **Explicación:** el objetivo plantea medir una variable, pero la metodología descrita no incluye un instrumento para hacerlo.
- **Sugerencia:** ajustar el instrumento o replantear el objetivo.

### Elemento ausente: Hipótesis
- No se identificó una hipótesis explícita en el texto proporcionado.

### Resumen
- **Nivel de coherencia general:** Medio (ejemplo)
- **Correcciones más urgentes:**
  1. Alinear objetivo y metodología
  2. Formular hipótesis explícita si el enfoque es cuantitativo
  3. Revisar que las conclusiones respondan directamente a los objetivos"""


def _armar_mensaje(idioma: str, nombre: str | None, texto: str) -> str:
    return f"""Idioma solicitado: {idioma}
Documento: {nombre or 'sin nombre'}

{texto[:20000]}"""


@router.post("/revisar", response_model=ReporteErrores)
def revisar_documento(
    solicitud: SolicitudRevision,
    usuario: str = Depends(usuario_autenticado),
) -> ReporteErrores:
    _validar_idioma(solicitud.idioma)
    try:
        texto = generar_respuesta(
            SYSTEM_PROMPT,
            _armar_mensaje(solicitud.idioma, solicitud.nombre_documento, solicitud.texto_documento),
            respuesta_simulada=_mock_reporte(solicitud.nombre_documento),
        )
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return ReporteErrores(reporte_markdown=texto)


@router.post("/revisar-con-archivo", response_model=ReporteErrores)
async def revisar_documento_con_archivo(
    idioma: str = Form(default="es"),
    nombre_documento: str | None = Form(default=None),
    texto_documento: str | None = Form(default=None),
    archivo: UploadFile | None = File(default=None),
    usuario: str = Depends(usuario_autenticado),
) -> ReporteErrores:
    _validar_idioma(idioma)

    texto_archivo = ""
    if archivo is not None:
        contenido = await archivo.read()
        tamano_mb = len(contenido) / (1024 * 1024)
        if tamano_mb > TAMANO_MAXIMO_ARCHIVO_MB:
            raise HTTPException(
                status_code=413,
                detail=(
                    f"El archivo pesa {tamano_mb:.1f} MB, y el máximo permitido "
                    f"es {TAMANO_MAXIMO_ARCHIVO_MB} MB."
                ),
            )
        if not archivo.filename.lower().endswith(".pdf"):
            raise HTTPException(status_code=422, detail="Solo se aceptan archivos PDF por ahora.")
        try:
            lector = PdfReader(io.BytesIO(contenido))
            texto_archivo = "\n".join(p.extract_text() or "" for p in lector.pages)
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"No se pudo leer el PDF: {e}")
        if not nombre_documento:
            nombre_documento = archivo.filename

    partes = [p for p in [texto_documento, texto_archivo] if p and p.strip()]
    texto_completo = "\n\n".join(partes)
    if not texto_completo.strip():
        raise HTTPException(
            status_code=422,
            detail="Debes pegar el texto del documento y/o adjuntar un PDF.",
        )

    try:
        texto = generar_respuesta(
            SYSTEM_PROMPT,
            _armar_mensaje(idioma, nombre_documento, texto_completo),
            respuesta_simulada=_mock_reporte(nombre_documento),
        )
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return ReporteErrores(reporte_markdown=texto)


# ---------- Exportar a Word ----------

class SolicitudExportar(BaseModel):
    reporte_markdown: str
    titulo: str = "Reporte de revisión metodológica"


def _agregar_texto_con_negritas(parrafo, texto: str) -> None:
    """Convierte **negritas** de Markdown en negritas reales de Word."""
    texto = texto.replace("<br>", "\n").replace("<br/>", "\n")
    for i, trozo in enumerate(texto.split("**")):
        if trozo:
            run = parrafo.add_run(trozo)
            run.bold = (i % 2 == 1)


def _es_fila_tabla(linea: str) -> bool:
    return linea.startswith("|") and linea.count("|") >= 2


@router.post("/exportar-word")
def exportar_word(
    solicitud: SolicitudExportar,
    usuario: str = Depends(usuario_autenticado),
) -> StreamingResponse:
    documento = DocxDocument()
    documento.add_heading(solicitud.titulo, level=1)

    lineas = solicitud.reporte_markdown.split("\n")
    i = 0
    while i < len(lineas):
        linea = lineas[i].strip()

        if not linea or re.fullmatch(r"-{3,}", linea):
            i += 1
            continue

        # Tablas Markdown: junta las filas consecutivas
        if _es_fila_tabla(linea):
            filas = []
            while i < len(lineas) and _es_fila_tabla(lineas[i].strip()):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in celdas if c):
                    filas.append(celdas)
                i += 1
            if filas:
                columnas = max(len(f) for f in filas)
                tabla = documento.add_table(rows=len(filas), cols=columnas)
                tabla.style = "Table Grid"
                for r, fila in enumerate(filas):
                    for c in range(columnas):
                        celda = tabla.cell(r, c)
                        celda.text = ""
                        _agregar_texto_con_negritas(
                            celda.paragraphs[0], fila[c] if c < len(fila) else ""
                        )
                documento.add_paragraph("")
            continue

        encabezado = re.match(r"^(#{1,4})\s+(.*)", linea)
        if encabezado:
            nivel = min(len(encabezado.group(1)) + 1, 4)
            documento.add_heading(encabezado.group(2).replace("**", ""), level=nivel)
        elif re.match(r"^[-*]\s+", linea):
            p = documento.add_paragraph(style="List Bullet")
            _agregar_texto_con_negritas(p, re.sub(r"^[-*]\s+", "", linea))
        elif re.match(r"^\d+[.)]\s+", linea):
            p = documento.add_paragraph(style="List Number")
            _agregar_texto_con_negritas(p, re.sub(r"^\d+[.)]\s+", "", linea))
        else:
            p = documento.add_paragraph()
            _agregar_texto_con_negritas(p, linea)
        i += 1

    buffer = io.BytesIO()
    documento.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="reporte_revision_metodologica.docx"'},
    )
