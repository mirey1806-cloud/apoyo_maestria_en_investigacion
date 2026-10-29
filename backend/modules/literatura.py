"""
Módulo 2: Analizador de literatura científica.

Sigue el mismo patrón que el Módulo 1 (protocolos.py):
- Requiere sesión iniciada (login)
- Soporta varios idiomas
- Permite pegar texto y/o adjuntar un PDF (o varios artículos separados por "---")
- Permite exportar el análisis generado a un archivo Word
"""
import io
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document as DocxDocument

from backend.core.llm_client import generar_respuesta
from backend.core.config import IDIOMAS_SOPORTADOS, TAMANO_MAXIMO_ARCHIVO_MB
from backend.modules.auth import usuario_autenticado

router = APIRouter(prefix="/api/literatura", tags=["Analizador de literatura"])

SYSTEM_PROMPT = """Eres un asistente de análisis de literatura científica para
estudiantes de posgrado. Recibirás el texto de uno o más artículos (si hay
varios, vienen separados por la línea "---").

Para CADA artículo entrega:
- Referencia aproximada (autor/año si aparece en el texto, o "no identificado")
- Objetivo del estudio
- Metodología empleada
- Hallazgos principales (máximo 5 puntos)
- Conceptos clave (lista corta)

Al final entrega una SÍNTESIS COMPARATIVA que identifique:
- Coincidencias entre los artículos
- Diferencias metodológicas o de resultados
- Vacíos de investigación (research gaps) que el estudiante podría explorar

Reglas:
- No inventes datos que no estén en el texto proporcionado.
- Deja explícito que esto es un insumo de trabajo, no un texto para copiar
  directamente en la tesis (riesgo de plagio si se usa sin reelaborar).
- Responde en el idioma que te indique el usuario.
"""


class SolicitudAnalisis(BaseModel):
    idioma: str = Field(default="es")
    articulos_texto: str = Field(..., description="Texto de uno o más artículos, separados por '---'")


class AnalisisLiteratura(BaseModel):
    analisis_markdown: str


def _validar_idioma(idioma: str) -> None:
    if idioma not in IDIOMAS_SOPORTADOS:
        raise HTTPException(
            status_code=422,
            detail=f"Idioma '{idioma}' no soportado. Usa uno de: {', '.join(IDIOMAS_SOPORTADOS)}",
        )


def _mock_analisis(articulos_texto: str) -> str:
    num_articulos = len([a for a in articulos_texto.split("---") if a.strip()]) or 1
    bloques = []
    for i in range(1, num_articulos + 1):
        bloques.append(f"""### Artículo {i}
Referencia: no identificada (resultado de ejemplo)
Objetivo del estudio: [se extraería del texto proporcionado]
Metodología: [se extraería del texto proporcionado]
Hallazgos principales:
1. Hallazgo de ejemplo 1
2. Hallazgo de ejemplo 2
Conceptos clave: [conceptos de ejemplo]""")
    bloques.append("""### Síntesis comparativa
Coincidencias: (ejemplo) los artículos comparten enfoque en el tema analizado.
Diferencias: (ejemplo) difieren en el diseño metodológico usado.
Vacíos de investigación detectados: (ejemplo) falta explorar la relación entre
las variables X e Y en el contexto local.

*Nota: este análisis es un insumo de trabajo; no debe copiarse tal cual en tu tesis.*""")
    return "\n\n".join(bloques)


@router.post("/generar", response_model=AnalisisLiteratura)
def generar_analisis(
    solicitud: SolicitudAnalisis,
    usuario: str = Depends(usuario_autenticado),
) -> AnalisisLiteratura:
    _validar_idioma(solicitud.idioma)

    mensaje = f"""Idioma solicitado: {solicitud.idioma}

Texto de él/los artículo(s):
{solicitud.articulos_texto[:20000]}"""

    try:
        texto = generar_respuesta(
            SYSTEM_PROMPT, mensaje,
            respuesta_simulada=_mock_analisis(solicitud.articulos_texto),
        )
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return AnalisisLiteratura(analisis_markdown=texto)


@router.post("/generar-con-archivo", response_model=AnalisisLiteratura)
async def generar_analisis_con_archivo(
    idioma: str = Form(default="es"),
    articulos_texto: str | None = Form(default=None),
    archivo: UploadFile | None = File(default=None),
    usuario: str = Depends(usuario_autenticado),
) -> AnalisisLiteratura:
    """
    Igual que /generar, pero además permite adjuntar un PDF con el artículo
    (o uno de los artículos) a analizar.
    """
    _validar_idioma(idioma)

    texto_archivo = ""
    if archivo is not None:
        contenido = await archivo.read()
        tamano_mb = len(contenido) / (1024 * 1024)
        if tamano_mb > TAMANO_MAXIMO_ARCHIVO_MB:
            raise HTTPException(
                status_code=413,
                detail=(
                    f"El archivo pesa {tamano_mb:.1f} MB, y el máximo permitido "
                    f"es {TAMANO_MAXIMO_ARCHIVO_MB} MB."
                ),
            )
        if not archivo.filename.lower().endswith(".pdf"):
            raise HTTPException(status_code=422, detail="Solo se aceptan archivos PDF por ahora.")
        try:
            lector = PdfReader(io.BytesIO(contenido))
            texto_archivo = "\n".join(pagina.extract_text() or "" for pagina in lector.pages)
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"No se pudo leer el PDF: {e}")

    partes = [p for p in [articulos_texto, texto_archivo] if p and p.strip()]
    texto_completo = "\n\n---\n\n".join(partes)

    if not texto_completo.strip():
        raise HTTPException(
            status_code=422,
            detail="Debes pegar texto y/o adjuntar un PDF con al menos un artículo.",
        )

    mensaje = f"""Idioma solicitado: {idioma}

Texto de él/los artículo(s):
{texto_completo[:20000]}"""

    try:
        texto = generar_respuesta(
            SYSTEM_PROMPT, mensaje,
            respuesta_simulada=_mock_analisis(texto_completo),
        )
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return AnalisisLiteratura(analisis_markdown=texto)


class SolicitudExportar(BaseModel):
    analisis_markdown: str
    titulo: str = "Análisis de literatura científica"


@router.post("/exportar-word")
def exportar_word(
    solicitud: SolicitudExportar,
    usuario: str = Depends(usuario_autenticado),
) -> StreamingResponse:
    documento = DocxDocument()
    documento.add_heading(solicitud.titulo, level=1)

    for linea in solicitud.analisis_markdown.split("\n"):
        linea_limpia = linea.strip()
        if not linea_limpia:
            documento.add_paragraph("")
            continue
        if linea_limpia.startswith("### "):
            documento.add_heading(linea_limpia.replace("### ", ""), level=2)
        elif linea_limpia.startswith(("- ", "* ")):
            documento.add_paragraph(linea_limpia[2:], style="List Bullet")
        elif linea_limpia[:2].isdigit() is False and linea_limpia[:1].isdigit() and ". " in linea_limpia[:4]:
            documento.add_paragraph(linea_limpia, style="List Number")
        else:
            documento.add_paragraph(linea_limpia)

    buffer = io.BytesIO()
    documento.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="analisis_literatura.docx"'},
    )
