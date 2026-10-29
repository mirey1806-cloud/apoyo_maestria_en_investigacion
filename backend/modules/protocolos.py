"""
Módulo 1: Constructor de protocolos de investigación.

Cambios en esta versión (según observaciones del asesor externo):
- Requiere sesión iniciada (login).
- Soporta más idiomas, no solo español/inglés.
- Permite adjuntar un PDF de contexto (además de escribirlo a mano),
  aceptando archivos más grandes que antes.
- Permite exportar el protocolo generado directamente a un archivo Word.
"""
import io
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document as DocxDocument

from backend.core.llm_client import generar_respuesta
from backend.core.config import IDIOMAS_SOPORTADOS, TAMANO_MAXIMO_ARCHIVO_MB
from backend.modules.auth import usuario_autenticado

router = APIRouter(prefix="/api/protocolos", tags=["Constructor de protocolos"])

PLANTILLA_PROTOCOLO = """
1. Título tentativo
2. Planteamiento del problema
3. Preguntas de investigación
4. Objetivo general
5. Objetivos específicos
6. Justificación
7. Marco teórico preliminar (principales conceptos y autores a revisar)
8. Enfoque y diseño metodológico (cualitativo / cuantitativo / mixto)
9. Población y muestra (si aplica)
10. Técnicas e instrumentos de recolección de datos
11. Cronograma tentativo (fases, no fechas exactas)
12. Referencias preliminares sugeridas (solo temas/autores, no inventar citas)
"""

SYSTEM_PROMPT = f"""Eres un asistente metodológico para estudiantes de posgrado en
investigación. Tu tarea es generar un BORRADOR de protocolo de investigación
siguiendo EXACTAMENTE esta plantilla, en el idioma que te indique el usuario:
{PLANTILLA_PROTOCOLO}

Reglas importantes:
- Este es un borrador de apoyo, no un documento final: debe quedar claro que
  requiere revisión del estudiante y su asesor.
- No inventes referencias bibliográficas reales ni datos falsos; solo sugiere
  líneas/temas a buscar.
- Sé específico al problema descrito por el usuario, no genérico.
- Usa un lenguaje académico claro y directo.
- Si el usuario adjuntó texto de contexto adicional (por ejemplo de un PDF),
  úsalo para enriquecer el protocolo, sin copiarlo textualmente.
"""


class DatosProyecto(BaseModel):
    tema: str = Field(..., description="Tema o área de investigación")
    problema: str = Field(..., description="Problema u oportunidad observada")
    idioma: str = Field(default="es")
    enfoque: str | None = Field(default=None)
    contexto: str | None = Field(default=None)


class ProtocoloGenerado(BaseModel):
    protocolo_markdown: str


def _validar_idioma(idioma: str) -> None:
    if idioma not in IDIOMAS_SOPORTADOS:
        raise HTTPException(
            status_code=422,
            detail=f"Idioma '{idioma}' no soportado. Usa uno de: {', '.join(IDIOMAS_SOPORTADOS)}",
        )


def _mock_protocolo(datos: DatosProyecto) -> str:
    return f"""1. Título tentativo
{datos.tema} — un estudio {datos.enfoque or '(enfoque a definir)'}

2. Planteamiento del problema
{datos.problema}

3. Preguntas de investigación
¿Cómo influye [variable relacionada con "{datos.tema}"] en el problema descrito?

4. Objetivo general
Analizar {datos.tema} en relación con el problema planteado.

5. Objetivos específicos
- Identificar los factores asociados
- Describir el comportamiento del fenómeno
- Proponer alternativas basadas en la evidencia recolectada

6. Justificación
Este estudio es relevante porque atiende una problemática observada en el
contexto descrito ({datos.contexto or 'contexto no especificado'}).

7. Marco teórico preliminar
Conceptos clave a revisar: [conceptos relacionados con {datos.tema}]

8. Enfoque y diseño metodológico
{datos.enfoque or 'A definir según el tipo de datos disponibles'}

9. Población y muestra
Por definir según el alcance del estudio.

10. Técnicas e instrumentos
Por definir (encuesta, entrevista, análisis documental, etc.)

11. Cronograma tentativo
Fase 1: Revisión de literatura — Fase 2: Recolección de datos — Fase 3: Análisis"""


@router.post("/generar", response_model=ProtocoloGenerado)
def generar_protocolo(
    datos: DatosProyecto,
    usuario: str = Depends(usuario_autenticado),
) -> ProtocoloGenerado:
    _validar_idioma(datos.idioma)

    mensaje = f"""Idioma solicitado: {datos.idioma}
Tema: {datos.tema}
Problema observado: {datos.problema}
Enfoque metodológico preferido: {datos.enfoque or 'no especificado, sugiere el más adecuado'}
Contexto adicional: {datos.contexto or 'ninguno'}

Genera el protocolo siguiendo la plantilla indicada."""

    try:
        texto = generar_respuesta(SYSTEM_PROMPT, mensaje, respuesta_simulada=_mock_protocolo(datos))
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return ProtocoloGenerado(protocolo_markdown=texto)


@router.post("/generar-con-archivo", response_model=ProtocoloGenerado)
async def generar_protocolo_con_archivo(
    tema: str = Form(...),
    problema: str = Form(...),
    idioma: str = Form(default="es"),
    enfoque: str | None = Form(default=None),
    contexto: str | None = Form(default=None),
    archivo: UploadFile | None = File(default=None),
    usuario: str = Depends(usuario_autenticado),
) -> ProtocoloGenerado:
    """
    Igual que /generar, pero además permite adjuntar un PDF (por ejemplo,
    un documento institucional o un avance previo) cuyo texto se extrae y
    se agrega como contexto adicional para la IA.
    """
    _validar_idioma(idioma)

    texto_archivo = ""
    if archivo is not None:
        contenido = await archivo.read()
        tamano_mb = len(contenido) / (1024 * 1024)
        if tamano_mb > TAMANO_MAXIMO_ARCHIVO_MB:
            raise HTTPException(
                status_code=413,
                detail=(
                    f"El archivo pesa {tamano_mb:.1f} MB, y el máximo permitido "
                    f"es {TAMANO_MAXIMO_ARCHIVO_MB} MB."
                ),
            )
        if not archivo.filename.lower().endswith(".pdf"):
            raise HTTPException(status_code=422, detail="Solo se aceptan archivos PDF por ahora.")
        try:
            lector = PdfReader(io.BytesIO(contenido))
            texto_archivo = "\n".join(pagina.extract_text() or "" for pagina in lector.pages)
            texto_archivo = texto_archivo[:20000]  # límite de contexto razonable
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"No se pudo leer el PDF: {e}")

    datos = DatosProyecto(
        tema=tema, problema=problema, idioma=idioma, enfoque=enfoque,
        contexto=(contexto or "") + ("\n\nContenido extraído del PDF adjunto:\n" + texto_archivo if texto_archivo else ""),
    )

    mensaje = f"""Idioma solicitado: {datos.idioma}
Tema: {datos.tema}
Problema observado: {datos.problema}
Enfoque metodológico preferido: {datos.enfoque or 'no especificado, sugiere el más adecuado'}
Contexto adicional (puede incluir texto extraído de un PDF adjunto): {datos.contexto or 'ninguno'}

Genera el protocolo siguiendo la plantilla indicada."""

    try:
        texto = generar_respuesta(SYSTEM_PROMPT, mensaje, respuesta_simulada=_mock_protocolo(datos))
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return ProtocoloGenerado(protocolo_markdown=texto)


class SolicitudExportar(BaseModel):
    protocolo_markdown: str
    titulo: str = "Protocolo de investigación"


@router.post("/exportar-word")
def exportar_word(
    solicitud: SolicitudExportar,
    usuario: str = Depends(usuario_autenticado),
) -> StreamingResponse:
    """
    Convierte el texto del protocolo generado (Markdown simple) en un
    archivo .docx descargable, para que el estudiante siga trabajándolo
    fuera de la plataforma.
    """
    documento = DocxDocument()
    documento.add_heading(solicitud.titulo, level=1)

    for linea in solicitud.protocolo_markdown.split("\n"):
        linea_limpia = linea.strip()
        if not linea_limpia:
            documento.add_paragraph("")
            continue
        if len(linea_limpia) > 2 and linea_limpia[0].isdigit() and ". " in linea_limpia[:4]:
            documento.add_heading(linea_limpia, level=2)
        elif linea_limpia.startswith(("- ", "* ")):
            documento.add_paragraph(linea_limpia[2:], style="List Bullet")
        else:
            documento.add_paragraph(linea_limpia)

    buffer = io.BytesIO()
    documento.save(buffer)
    buffer.seek(0)

    nombre_archivo = "protocolo_investigacion.docx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )
